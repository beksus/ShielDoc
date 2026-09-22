import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app.analyzer import analyze_pages
from app.extractors.docx import extract_docx


class DOCXPipelineTests(unittest.TestCase):
    def test_extracts_and_detects_email_in_paragraph(self):
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "synthetic.docx"

            document = Document()
            document.add_paragraph("No email in this paragraph.")
            document.add_paragraph("Contact: demo@example.com")
            document.save(str(docx_path))

            paragraphs = extract_docx(str(docx_path))
            entities = analyze_pages(paragraphs)

            self.assertEqual(len(paragraphs), 2)
            self.assertEqual(len(entities), 1)

            entity = entities[0]

            self.assertEqual(entity.type, "EMAIL")
            self.assertEqual(entity.value, "demo@example.com")
            self.assertIsNone(entity.page)

            paragraph_text = paragraphs[1]["text"]
            self.assertEqual(
                paragraph_text[entity.start:entity.end],
                entity.value,
            )
            
    def test_detects_email_and_phone_from_docx(self):
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "synthetic_contacts.docx"

            document = Document()
            document.add_paragraph("Email: demo@example.com")
            document.add_paragraph("Phone: +7 (000) 000-00-00")
            document.save(str(docx_path))

            paragraphs = extract_docx(str(docx_path))
            entities = analyze_pages(paragraphs)

            self.assertEqual(
                [(entity.type, entity.value) for entity in entities],
                [
                    ("EMAIL", "demo@example.com"),
                    ("PHONE", "+7 (000) 000-00-00"),
                ],
            )

            email, phone = entities

            self.assertIsNone(email.page)
            self.assertIsNone(phone.page)

            self.assertEqual(
                paragraphs[0]["text"][email.start:email.end],
                email.value,
            )
            self.assertEqual(
                paragraphs[1]["text"][phone.start:phone.end],
                phone.value,
            )

    def test_preserves_paragraph_numbers(self):
        with TemporaryDirectory() as directory:
            docx_path = Path(directory) / "synthetic.docx"

            document = Document()
            document.add_paragraph("Email: demo@example.com")
            document.add_paragraph("")
            document.add_paragraph("Email: demo@example.com")
            document.save(str(docx_path))

            paragraphs = extract_docx(str(docx_path))
            entities = analyze_pages(paragraphs)

            self.assertEqual(len(entities), 2)
            self.assertEqual(
                [entity.paragraph for entity in entities],
                [1, 3],
            )

            for entity in entities:
                self.assertIsNone(entity.page)

                text = paragraphs[entity.paragraph - 1]["text"]
                self.assertEqual(
                    text[entity.start:entity.end],
                    entity.value,
                )