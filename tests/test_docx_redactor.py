import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app.analyzer import analyze_pages
from app.extractors.docx import extract_docx
from app.redactors.docx import redact_docx


class DOCXRedactorTests(unittest.TestCase):
    def test_redacts_copy_and_preserves_original(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic.docx"
            output = Path(directory) / "redacted.docx"

            document = Document()
            document.add_paragraph("Email: demo@example.com")
            document.add_paragraph("")
            document.add_paragraph("Email: demo@example.com")
            document.add_paragraph("Keep this paragraph.")
            document.save(str(source))

            original_bytes = source.read_bytes()
            entities = analyze_pages(extract_docx(str(source)))

            redact_docx(source, output, entities)

            saved = Document(str(output))

            self.assertEqual(
                [paragraph.text for paragraph in saved.paragraphs],
                [
                    "Email: [REDACTED]",
                    "",
                    "Email: [REDACTED]",
                    "Keep this paragraph.",
                ],
            )
            self.assertEqual(source.read_bytes(), original_bytes)
            remaining_entities = analyze_pages(
                extract_docx(str(output))
            )
            self.assertEqual(remaining_entities, [])

    def test_refuses_to_overwrite_source(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic.docx"
            Document().save(str(source))

            original_bytes = source.read_bytes()

            with self.assertRaises(ValueError):
                redact_docx(source, source, [])

            self.assertEqual(source.read_bytes(), original_bytes)

    def test_refuses_to_overwrite_existing_output(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / "synthetic.docx"
            output = Path(directory) / "existing.docx"

            Document().save(str(source))
            Document().save(str(output))
            original_bytes = output.read_bytes()

            with self.assertRaises(FileExistsError):
                redact_docx(source, output, [])

            self.assertEqual(output.read_bytes(), original_bytes)